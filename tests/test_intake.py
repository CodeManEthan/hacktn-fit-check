"""Reading a clinic's wish-list, from a spreadsheet or from typed text.

No test here touches the network: the free-text parser is always handed a
stand-in client.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Sequence

import pytest
from openpyxl import Workbook

from app.ai.client import GenAiError
from app.intake.freetext import MAX_ITEMS, parse_list
from app.intake.tables import MAX_ROWS, IntakeError, parse_table


def csv_bytes(*rows: str) -> bytes:
    """A .csv upload, written out exactly as the clinic's file would be."""
    return "\n".join(rows).encode("utf-8")


def xlsx_bytes(rows: Sequence[Sequence[Any]]) -> bytes:
    """A real .xlsx upload, built in memory so no fixture file is needed."""
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(list(row))
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class StubClient:
    """Stands in for the OCI client. Returns a payload, or raises on demand."""

    def __init__(
        self, payload: dict[str, Any] | None = None, error: Exception | None = None
    ) -> None:
        self._payload = payload if payload is not None else {}
        self._error = error
        self.calls: list[tuple[str, str]] = []

    def complete_json(self, prompt: str, *, system: str = "", **_: Any) -> dict[str, Any]:
        self.calls.append((prompt, system))
        if self._error is not None:
            raise self._error
        return self._payload


def items(*pairs: tuple[str, int]) -> dict[str, Any]:
    """The shape the model is asked to answer in."""
    return {"items": [{"device_type": name, "quantity": count} for name, count in pairs]}


# --- the package door -------------------------------------------------------


def test_both_parsers_are_reachable_from_the_package():
    # Arrange / Act
    import app.intake as intake

    # Assert
    assert intake.parse_table is parse_table
    assert intake.parse_list is parse_list
    assert intake.IntakeError is IntakeError


# --- spreadsheets -----------------------------------------------------------


def test_a_csv_with_headers_becomes_one_line_per_row():
    # Arrange
    raw = csv_bytes("Device,Quantity", "Ventilator,2", "Oxygen concentrator,3")

    # Act
    lines = parse_table(raw, "wishlist.csv")

    # Assert
    assert [(line.device_type, line.quantity) for line in lines] == [
        ("ventilator", 2),
        ("oxygen concentrator", 3),
    ]


def test_headers_are_found_whatever_their_case_and_spacing():
    # Arrange: the clinic typed its own headers, loosely.
    raw = csv_bytes("  ITEM  ,  Qty ", "Infusion pump,4")

    # Act
    lines = parse_table(raw, "wishlist.csv")

    # Assert
    assert lines[0].device_type == "infusion pump"
    assert lines[0].quantity == 4


def test_a_csv_without_headers_reads_the_first_two_columns():
    # Arrange: no header row at all, just items and counts.
    raw = csv_bytes("Ventilator,2", "Patient monitor,1")

    # Act
    lines = parse_table(raw, "wishlist.csv")

    # Assert
    assert [(line.device_type, line.quantity) for line in lines] == [
        ("ventilator", 2),
        ("patient monitor", 1),
    ]


def test_a_headerless_second_column_that_is_not_a_number_is_ignored():
    # Arrange: the second column holds a note, not a count.
    raw = csv_bytes("Ventilator,urgent", "Patient monitor,spare")

    # Act
    lines = parse_table(raw, "wishlist.csv")

    # Assert
    assert [(line.device_type, line.quantity) for line in lines] == [
        ("ventilator", 1),
        ("patient monitor", 1),
    ]


def test_an_xlsx_with_headers_becomes_request_lines():
    # Arrange
    raw = xlsx_bytes([["Equipment", "Count"], ["Ultrasound", 2], ["Defibrillator", 1]])

    # Act
    lines = parse_table(raw, "wishlist.xlsx")

    # Assert
    assert [(line.device_type, line.quantity) for line in lines] == [
        ("ultrasound", 2),
        ("defibrillator", 1),
    ]


def test_a_missing_quantity_column_asks_for_one_of_each():
    # Arrange: a single named column, no counts anywhere.
    raw = csv_bytes("Device", "Ventilator", "Ultrasound")

    # Act
    lines = parse_table(raw, "wishlist.csv")

    # Assert
    assert [line.quantity for line in lines] == [1, 1]


def test_a_blank_quantity_cell_asks_for_one_unit():
    # Arrange: the clinic filled in the count for one row only.
    raw = csv_bytes("Device,Quantity", "Ventilator,", "Ultrasound,3")

    # Act
    lines = parse_table(raw, "wishlist.csv")

    # Assert
    assert [(line.device_type, line.quantity) for line in lines] == [
        ("ventilator", 1),
        ("ultrasound", 3),
    ]


def test_an_unparseable_quantity_asks_for_one_unit():
    # Arrange: "a few" is imprecise, not a reason to reject the upload.
    raw = csv_bytes("Device,Quantity", "Ventilator,a few", "Ultrasound,-4")

    # Act
    lines = parse_table(raw, "wishlist.csv")

    # Assert
    assert [line.quantity for line in lines] == [1, 1]


def test_blank_rows_and_nameless_rows_are_skipped():
    # Arrange: trailing blank lines and a count with nothing to count.
    raw = csv_bytes("Device,Quantity", "Ventilator,2", "", "   ,5", ",", "Ultrasound,1")

    # Act
    lines = parse_table(raw, "wishlist.csv")

    # Assert
    assert [line.device_type for line in lines] == ["ventilator", "ultrasound"]


def test_blank_rows_in_a_spreadsheet_are_skipped():
    # Arrange: openpyxl hands back a row of Nones for an untouched row.
    raw = xlsx_bytes([["Device", "Quantity"], ["Ventilator", 2], [None, None], ["Ultrasound", 1]])

    # Act
    lines = parse_table(raw, "wishlist.xlsx")

    # Assert
    assert [line.device_type for line in lines] == ["ventilator", "ultrasound"]


def test_the_same_device_asked_for_twice_stays_two_lines():
    # Arrange: two wards each asked for a ventilator; neither is dropped.
    raw = csv_bytes("Device,Quantity", "Ventilator,2", "Ventilator,3")

    # Act
    lines = parse_table(raw, "wishlist.csv")

    # Assert
    assert len(lines) == 2
    assert [line.quantity for line in lines] == [2, 3]


def test_a_long_upload_is_capped():
    # Arrange: far more rows than any clinic would file by hand.
    rows = ["Device,Quantity"] + [f"Device {index},1" for index in range(MAX_ROWS + 50)]

    # Act
    lines = parse_table(csv_bytes(*rows), "wishlist.csv")

    # Assert
    assert len(lines) == MAX_ROWS


def test_an_unsupported_extension_is_refused():
    # Arrange / Act / Assert
    with pytest.raises(IntakeError) as raised:
        parse_table(b"Ventilator,2", "wishlist.pdf")

    assert "wishlist.pdf" in str(raised.value)


def test_an_empty_file_is_refused():
    with pytest.raises(IntakeError):
        parse_table(b"", "wishlist.csv")


def test_a_file_with_only_headers_is_refused():
    # Arrange: a template nobody filled in.
    with pytest.raises(IntakeError):
        parse_table(csv_bytes("Device,Quantity"), "wishlist.csv")


def test_headers_below_a_blank_first_row_are_still_found():
    # Arrange: exports often carry an empty first line.
    raw = csv_bytes("", "Device,Quantity", "Ventilator,2")

    # Act
    lines = parse_table(raw, "wishlist.csv")

    # Assert
    assert [(line.device_type, line.quantity) for line in lines] == [("ventilator", 2)]


def test_bytes_that_are_not_a_spreadsheet_are_refused():
    with pytest.raises(IntakeError):
        parse_table(b"not a workbook at all", "wishlist.xlsx")


def test_a_workbook_with_no_data_sheet_is_refused():
    # Arrange: a valid .xlsx holding only a chart, so there are no rows to read.
    workbook = Workbook()
    workbook.create_chartsheet("Chart")
    workbook.remove(workbook["Sheet"])
    stream = BytesIO()
    workbook.save(stream)

    # Act / Assert
    with pytest.raises(IntakeError):
        parse_table(stream.getvalue(), "wishlist.xlsx")


def test_a_csv_cell_too_large_to_read_is_refused():
    # Arrange: one field past the csv module's limit, as a hostile upload would be.
    raw = b'Device,Quantity\n"' + b"a" * 200_000

    # Act / Assert
    with pytest.raises(IntakeError):
        parse_table(raw, "wishlist.csv")


def test_a_csv_that_is_not_text_is_refused():
    with pytest.raises(IntakeError):
        parse_table(b"\xff\xfe\x81\x00device", "wishlist.csv")


# --- free text --------------------------------------------------------------


def test_the_model_reading_of_a_typed_list_becomes_request_lines():
    # Arrange
    client = StubClient(items(("ultrasound", 2), ("x-ray machine", 1)))

    # Act
    lines = parse_list("2 ultrasounds and an x-ray machine", client=client)

    # Assert
    assert [(line.device_type, line.quantity) for line in lines] == [
        ("ultrasound", 2),
        ("x-ray machine", 1),
    ]


def test_the_typed_text_and_a_system_prompt_are_sent_to_the_model():
    # Arrange
    client = StubClient(items(("ventilator", 1)))

    # Act
    parse_list("one ventilator", client=client)

    # Assert
    prompt, system = client.calls[0]
    assert "one ventilator" in prompt
    assert system.strip()


def test_a_typed_list_still_parses_when_the_model_is_unreachable():
    # Arrange: the fallback has to stand in for the model entirely.
    client = StubClient(error=GenAiError("OCI is down"))

    # Act
    lines = parse_list("2 ultrasounds and an x-ray machine", client=client)

    # Assert
    assert [(line.device_type, line.quantity) for line in lines] == [
        ("ultrasound", 2),
        ("x-ray machine", 1),
    ]


def test_the_fallback_splits_a_newline_list_and_reads_leading_counts():
    # Arrange
    client = StubClient(error=GenAiError("OCI is down"))

    # Act
    lines = parse_list("x3 ventilators\n- Oxygen concentrator\n2 infusion pumps", client=client)

    # Assert
    assert [(line.device_type, line.quantity) for line in lines] == [
        ("ventilator", 3),
        ("oxygen concentrator", 1),
        ("infusion pump", 2),
    ]


def test_the_fallback_splits_a_comma_separated_list():
    # Arrange
    client = StubClient(error=GenAiError("OCI is down"))

    # Act
    lines = parse_list("ventilator, patient monitor, suction pump", client=client)

    # Assert
    assert [line.device_type for line in lines] == [
        "ventilator",
        "patient monitor",
        "suction pump",
    ]


def test_a_reply_that_is_not_the_agreed_shape_falls_back_to_splitting():
    # Arrange: valid JSON, wrong contents.
    client = StubClient({"result": "I cannot help with that"})

    # Act
    lines = parse_list("ventilator, ultrasound", client=client)

    # Assert
    assert [line.device_type for line in lines] == ["ventilator", "ultrasound"]


def test_nameless_items_in_the_model_reply_are_dropped():
    # Arrange
    client = StubClient({"items": [{"device_type": "  ", "quantity": 2}, "ventilator"]})

    # Act
    lines = parse_list("ventilator", client=client)

    # Assert: nothing usable came back, so the typed text is split instead.
    assert [line.device_type for line in lines] == ["ventilator"]


def test_a_count_the_model_got_wrong_becomes_one_unit():
    # Arrange: a request line may never ask for fewer than one unit.
    client = StubClient({"items": [{"device_type": "ventilator", "quantity": 0}]})

    # Act
    lines = parse_list("ventilators", client=client)

    # Assert
    assert lines[0].quantity == 1


def test_a_very_long_typed_list_is_capped():
    # Arrange
    client = StubClient(error=GenAiError("OCI is down"))
    text = ", ".join(f"device {index}" for index in range(MAX_ITEMS + 20))

    # Act
    lines = parse_list(text, client=client)

    # Assert
    assert len(lines) == MAX_ITEMS


def test_the_model_reply_is_capped_too():
    # Arrange
    client = StubClient(items(*[(f"device {index}", 1) for index in range(MAX_ITEMS + 20)]))

    # Act
    lines = parse_list("a long list", client=client)

    # Assert
    assert len(lines) == MAX_ITEMS


def test_the_fallback_drops_a_fragment_that_names_no_equipment():
    # Arrange: a stray count and an empty gap between commas.
    client = StubClient(error=GenAiError("OCI is down"))

    # Act
    lines = parse_list("2, , ventilator", client=client)

    # Assert
    assert [(line.device_type, line.quantity) for line in lines] == [("ventilator", 1)]


@pytest.mark.parametrize("typed,expected", [("2 apparatus", "apparatus"), ("2 gauze", "gauze")])
def test_the_fallback_leaves_a_word_that_is_not_a_plural_alone(typed: str, expected: str):
    # Arrange
    client = StubClient(error=GenAiError("OCI is down"))

    # Act
    lines = parse_list(typed, client=client)

    # Assert
    assert lines[0].device_type == expected


def test_the_terms_the_platform_stocks_are_offered_to_the_model():
    # Arrange: the clinic's wording has to land on something the inventory holds.
    client = StubClient(items(("x-ray system", 1)))

    # Act
    lines = parse_list("an x-ray machine", client=client, vocabulary=["x-ray system", "ventilator"])

    # Assert
    prompt, _ = client.calls[0]
    assert "x-ray system" in prompt
    assert lines[0].device_type == "x-ray system"


def test_empty_text_asks_for_nothing_and_never_calls_the_model():
    # Arrange
    client = StubClient(items(("ventilator", 1)))

    # Act
    lines = parse_list("   \n  ", client=client)

    # Assert
    assert lines == ()
    assert client.calls == []


# --- [B14] the routes in this tree: upload, standing pages, 404, clinic form --
#
# Each test gets its own SQLite file and a TestClient without the lifespan, so
# no seed runs. No model is called: the typed path is given a StubClient.

from dataclasses import replace as _replace

from fastapi.testclient import TestClient

from app.checks.runner import ALL_CHECKS, ROUTE_CHECKS
from app.config import SETTINGS
from app.data import db, repository
from app.data import marketplace_repo as repo
from app.domain.marketplace import GeoPoint, Offer, OfferStatus
from app.main import app
from app.reference.country_power import known_countries
from app.web import pages_routes
from tests.factories import device, facility


@pytest.fixture
def web(tmp_path, monkeypatch):
    url = f"sqlite:///{tmp_path / 'intake.db'}"
    monkeypatch.setattr(db, "SETTINGS", _replace(SETTINGS, database_url=url))
    monkeypatch.setattr(db, "_engine", None)
    repository.setup()
    repo.save_facility(facility())
    yield TestClient(app)
    if db._engine is not None:
        db._engine.dispose()


def _upload(web, **data):
    files = data.pop("files", None)
    return web.post("/request/upload", data=data, files=files, follow_redirects=False)


def test_an_uploaded_csv_files_a_request_and_lands_on_its_matches_page(web):
    reply = _upload(web, facility_slug="test-clinic", files={
        "list_file": ("wants.csv", csv_bytes("Device,Quantity", "oxygen concentrator,2"),
                      "text/csv")})

    assert reply.status_code == 303
    location = reply.headers["location"]
    assert location.startswith("/request/") and location.endswith("/matches")
    stored = repo.load_request(int(location.split("/")[2]))
    assert [(l.device_type, l.quantity) for l in stored.lines] == [("oxygen concentrator", 2)]
    assert web.get(location).status_code == 200


def test_a_typed_list_is_read_with_the_stocked_vocabulary(web, monkeypatch):
    stub = StubClient(items(("oxygen concentrator", 3)))
    seen = {}

    def parse(text, *, vocabulary=()):
        seen["vocabulary"] = list(vocabulary)
        return parse_list(text, client=stub, vocabulary=vocabulary)

    monkeypatch.setattr(pages_routes, "parse_list", parse)
    repo.save_offer(Offer(
        donor_name="Mercy", donor_email="m@example.org", device=device(),
        device_type="oxygen concentrator",
        origin=GeoPoint(country="Testland", city="Here", latitude=0.1, longitude=0.1),
        quantity_offered=1))

    reply = _upload(web, facility_slug="test-clinic", wanted_text="three concentrators")

    assert reply.status_code == 303 and reply.headers["location"].endswith("/matches")
    assert seen["vocabulary"] == ["oxygen concentrator"]
    stored = repo.load_request(int(reply.headers["location"].split("/")[2]))
    assert [(l.device_type, l.quantity) for l in stored.lines] == [("oxygen concentrator", 3)]


def test_an_upload_with_neither_file_nor_text_is_refused(web):
    reply = _upload(web, facility_slug="test-clinic")
    assert reply.status_code == 400 and "Upload a list" in reply.text


def test_an_upload_for_an_unregistered_facility_is_refused(web):
    reply = _upload(web, facility_slug="nowhere", wanted_text="ventilator")
    assert reply.status_code == 400 and "not registered" in reply.text


def test_a_bad_spreadsheet_says_why_on_the_error_page(web):
    reply = _upload(web, facility_slug="test-clinic", files={
        "list_file": ("wants.pdf", b"%PDF", "application/pdf")})
    assert reply.status_code == 400 and "not a spreadsheet" in reply.text


def test_the_upload_form_lists_the_registered_facilities(web):
    reply = web.get("/request/upload")
    assert reply.status_code == 200 and 'value="test-clinic"' in reply.text


def test_known_device_types_lists_open_offers_only(web):
    base = dict(donor_name="Mercy", donor_email="m@example.org", device=device(),
                origin=GeoPoint(country="Testland", city="Here", latitude=0.1, longitude=0.1),
                quantity_offered=1)
    repo.save_offer(Offer(device_type="suction pump", **base))
    repo.save_offer(Offer(device_type="ventilator", status=OfferStatus.WITHDRAWN, **base))
    assert repo.known_device_types() == ["suction pump"]


def test_the_about_page_states_the_check_count_from_the_registry(web):
    reply = web.get("/about")
    count = len(ALL_CHECKS) + len(ROUTE_CHECKS)
    assert reply.status_code == 200
    assert f"{count} checks written" in reply.text
    assert "reads the facts off it" not in reply.text
    assert "Import rules" in reply.text and "Uganda" in reply.text


def test_help_and_contact_render_without_substitutes_or_placeholder_mail(web):
    help_page = web.get("/help")
    contact = web.get("/contact")
    assert help_page.status_code == 200 and contact.status_code == 200
    assert "substitute" not in help_page.text.lower()
    assert "fit-check.example" not in help_page.text + contact.text
    assert "HackTN table" in contact.text


def test_an_unknown_address_renders_the_not_found_page(web):
    reply = web.get("/no/such/page")
    assert reply.status_code == 404
    assert "That page does not exist" in reply.text and "/no/such/page" in reply.text


def test_the_clinic_form_offers_the_country_list_and_no_three_phase_box(web):
    reply = web.get("/clinic")
    assert reply.status_code == 200
    assert '<select id="f-country" name="country">' in reply.text
    for name in ("Uganda", "Haiti", "Nepal"):
        assert name in known_countries() and f'<option value="{name}">' in reply.text
    assert "has_three_phase_power" not in reply.text


def test_registering_a_clinic_goes_to_the_list_upload(web):
    reply = web.post("/clinic", data={"name": "New Clinic", "country": "Uganda"},
                     follow_redirects=False)
    assert reply.status_code == 303
    assert reply.headers["location"] == "/request/upload?facility=new-clinic"
    page = web.get(reply.headers["location"])
    assert '<option value="new-clinic" selected>' in page.text
